# -*- coding: utf-8 -*-
# 观察池 Gordon 理论回报测算：单 sheet 明细（计算列=Excel公式带除零保护）+ 口径说明
import json

try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

def xl_color(css_hex):
    value = css_hex.removeprefix("#").upper()
    if len(value) != 6:
        raise ValueError(f"Expected #RRGGBB, got: {css_hex}")
    return "FF" + value

XL_PRIMARY = xl_color("#4472C4")
XL_HEAD_FONT = xl_color("#FFFFFF")
XL_INPUT = xl_color("#D9E2F3")
XL_CALC = xl_color("#F2F2F2")
XL_BORDER = xl_color("#BFBFBF")
XL_GREEN_BG = xl_color("#C6EFCE")
XL_GREEN_FT = xl_color("#006100")
XL_RED_BG = xl_color("#FFC7CE")
XL_RED_FT = xl_color("#9C0006")

thin = Side(style="thin", color=XL_BORDER)
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

rows = json.load(open("/tmp/gordon_sheet.json"))
N = len(rows)              # 31
TITLE_ROW = 1
HEAD_ROW = 2
DATA_START = 3
DATA_END = DATA_START + N - 1   # 33
AVG_ROW = DATA_END + 1          # 34

wb = openpyxl.Workbook()
wb.properties.title = "观察池Gordon理论回报测算"

# ══ Sheet1 测算明细 ══
ws = wb.active
ws.title = "测算明细"

ws.cell(row=TITLE_ROW, column=1, value="观察池 Gordon 理论回报测算（回报 = 股息率 + 内生增长 g；数据 2026-10-16 收盘）")
ws.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=13)
ws.cell(row=TITLE_ROW, column=1).font = Font(bold=True, size=13)
ws.cell(row=TITLE_ROW, column=1).alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[TITLE_ROW].height = 24

headers = ["序号", "标的", "市场", "现价", "PE(TTM)", "PB", "ROE", "股息率", "分红率",
           "内生增长g", "理论回报", "状态", "备注"]
for ci, h in enumerate(headers, start=1):
    c = ws.cell(row=HEAD_ROW, column=ci, value=h)
    c.font = Font(bold=True, color=XL_HEAD_FONT)
    c.fill = PatternFill("solid", fgColor=XL_PRIMARY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = border_all

pct2 = "0.00%"
pct1 = "0.0%"
for i, r in enumerate(rows):
    row = DATA_START + i
    vals = {
        1: (r["no"], "0", "center"),
        2: (r["name"], None, "left"),
        3: (r["mkt"], None, "center"),
        4: (r["px"], "0.00", "right"),
        5: (r["pe"], "0.0", "right"),
        6: (r["pb"], "0.00", "right"),
        7: (f'=IF(OR(E{row}="",E{row}=0,F{row}=""),"",F{row}/E{row})', pct1, "right"),          # ROE=PB/PE
        8: (r["dy"], pct2, "right"),                                                             # 股息率(输入)
        9: (f'=IF(OR(E{row}="",E{row}=0),"",H{row}*E{row})', pct1, "right"),                     # 分红率=股息率×PE
        10: (f'=IF(OR(G{row}="",G{row}=0),"",G{row}*(1-MIN(I{row},1)))', pct1, "right"),         # g=ROE×(1-截断分红率)
        11: (f'=IF(OR(H{row}="",J{row}=""),"",H{row}+J{row})', pct1, "right"),                   # 理论回报
        12: ("● 持仓" if r["held"] else "观察", None, "center"),
        13: (r["note"], None, "left"),
    }
    for ci, (v, fmt, align) in vals.items():
        c = ws.cell(row=row, column=ci, value=v)
        c.border = border_all
        if fmt:
            c.number_format = fmt
        c.alignment = Alignment(horizontal=align, vertical="center", wrapText=(ci == 13))
        if ci in (5, 6, 8):
            c.fill = PatternFill("solid", fgColor=XL_INPUT)   # 输入列浅蓝底
        elif ci in (7, 9, 10, 11):
            c.fill = PatternFill("solid", fgColor=XL_CALC)    # 计算列浅灰底

# 平均行
ws.cell(row=AVG_ROW, column=1, value="—")
ws.cell(row=AVG_ROW, column=2, value="观察池平均（简单算术）")
ws.cell(row=AVG_ROW, column=7, value=f'=IFERROR(AVERAGE(G{DATA_START}:G{DATA_END}),"")').number_format = pct1
ws.cell(row=AVG_ROW, column=8, value=f'=IFERROR(AVERAGE(H{DATA_START}:H{DATA_END}),"")').number_format = pct2
ws.cell(row=AVG_ROW, column=9, value=f'=IFERROR(AVERAGE(I{DATA_START}:I{DATA_END}),"")').number_format = pct1
ws.cell(row=AVG_ROW, column=10, value=f'=IFERROR(AVERAGE(J{DATA_START}:J{DATA_END}),"")').number_format = pct1
ws.cell(row=AVG_ROW, column=11, value=f'=IFERROR(AVERAGE(K{DATA_START}:K{DATA_END}),"")').number_format = pct1
for ci in range(1, 14):
    c = ws.cell(row=AVG_ROW, column=ci)
    c.font = Font(bold=True)
    c.border = border_all
    c.fill = PatternFill("solid", fgColor=xl_color("#D9E2F3"))

# 条件格式：理论回报 >=15% 绿、<8% 红
rng = f"K{DATA_START}:K{DATA_END}"
ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["0.15"],
    fill=PatternFill("solid", fgColor=XL_GREEN_BG), font=Font(color=XL_GREEN_FT, bold=True)))
ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0.08"],
    fill=PatternFill("solid", fgColor=XL_RED_BG), font=Font(color=XL_RED_FT)))

widths = {"A": 6, "B": 13, "C": 7, "D": 9.5, "E": 9, "F": 7.5, "G": 8.5, "H": 9, "I": 8.5, "J": 11, "K": 10.5, "L": 8.5, "M": 42}
for col, w in widths.items():
    ws.column_dimensions[col].width = w
ws.freeze_panes = "D3"
ws.auto_filter.ref = f"A{HEAD_ROW}:M{DATA_END}"

# ══ Sheet2 口径说明 ══
ws2 = wb.create_sheet("口径说明")
notes = [
    ("公式", "理论回报 = 股息率 + 内生增长g；g = ROE × (1-分红率)；分红率 = 股息率 × PE（>100%按100%截断）；ROE = PB ÷ PE(TTM)（同源恒等式）"),
    ("数据源", "腾讯行情源 gtimg（2026-10-16 收盘）：现价、PE(TTM)、PB、股息率（A股为行情字段直取，港股=每股DPS÷现价）"),
    ("覆盖", "投资平台观察池 31 只标的（含 7 只已持仓），无遗漏"),
    ("理论前提", "① PB 买卖时点不变；② ROE 可持续；③ 留存利润能赚出与存量相同的 ROE。任一崩塌则实际回报低于理论值——价值陷阱的来源"),
    ("† 泡泡玛特", "TTM ROE 53.2% 为利润率跳升年份的冲高值，g=39.5% 是把冲高值当永续外推的产物，实际可持续回报远低于 41.4%"),
    ("? 携程-S", "PE 5.4 疑似行情源字段错位（真实约 17-18），整行存疑待核"),
    ("‡ A/H 双上市", "H 股 ROE 受本币折算影响，与 A 股有 1-3 个点噪声（如平安 A 15.8% vs H 12.7%），H 股数字仅供参考"),
    ("对照", "书例（10% ROE/50%分红/2倍PB）= 7.5% 回报，约为本观察池中位水平；持仓 7 只全部在 11.7%-16.7% 区间"),
]
ws2.cell(row=1, column=1, value="口径说明").font = Font(bold=True, size=13)
for i, (k, v) in enumerate(notes, start=3):
    a = ws2.cell(row=i, column=1, value=k)
    a.font = Font(bold=True)
    a.alignment = Alignment(vertical="top")
    b = ws2.cell(row=i, column=2, value=v)
    b.alignment = Alignment(wrapText=True, vertical="top")
    ws2.row_dimensions[i].height = 30
ws2.column_dimensions["A"].width = 12
ws2.column_dimensions["B"].width = 100

out = "/Users/jimislg/Documents/郎刚的芒格/codex-产出/观察池Gordon理论回报测算-20261017.xlsx"
wb.save(out)
print("saved:", out, "| sheets:", wb.sheetnames)
