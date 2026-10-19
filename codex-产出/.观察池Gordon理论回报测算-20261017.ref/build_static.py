# -*- coding: utf-8 -*-
# 静态快照版：计算列由脚本算好写入数值（公式口径记录在口径说明页）
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule

def xl_color(css_hex):
    v = css_hex.removeprefix("#").upper()
    return "FF" + v

rows = json.load(open("/tmp/gordon_sheet.json"))
raw = {x["code"]: x for x in json.load(open("/tmp/gordon_rows.json"))}
N = len(rows)
DATA_START = 3
DATA_END = DATA_START + N - 1
AVG_ROW = DATA_END + 1

wb = openpyxl.Workbook()
wb.properties.title = "观察池Gordon理论回报测算"
ws = wb.active
ws.title = "测算明细"

ws.cell(row=1, column=1, value="观察池 Gordon 理论回报测算（回报 = 股息率 + 内生增长 g；数据 2026-10-16 收盘）")
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=13)
ws.cell(row=1, column=1).font = Font(bold=True, size=13)
ws.cell(row=1, column=1).alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[1].height = 24

headers = ["序号", "标的", "市场", "现价", "PE(TTM)", "PB", "ROE", "股息率", "分红率", "内生增长g", "理论回报", "状态", "备注"]
for ci, h in enumerate(headers, 1):
    c = ws.cell(row=2, column=ci, value=h)
    c.font = Font(bold=True, color=xl_color("#FFFFFF"))
    c.fill = PatternFill("solid", fgColor=xl_color("#4472C4"))
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = Border(left=Side("thin", color=xl_color("#BFBFBF")), right=Side("thin", color=xl_color("#BFBFBF")), top=Side("thin", color=xl_color("#BFBFBF")), bottom=Side("thin", color=xl_color("#BFBFBF")))

pct2, pct1 = "0.00%", "0.0%"
thinb = Border(left=Side("thin", color=xl_color("#BFBFBF")), right=Side("thin", color=xl_color("#BFBFBF")), top=Side("thin", color=xl_color("#BFBFBF")), bottom=Side("thin", color=xl_color("#BFBFBF")))
for i, r in enumerate(rows):
    row = DATA_START + i
    src = raw[r["code"]]
    payout = min(src["payout"], 1.0)
    vals = {
        1: (r["no"], "0", "center"), 2: (r["name"], None, "left"), 3: (r["mkt"], None, "center"),
        4: (r["px"], "0.00", "right"), 5: (r["pe"], "0.0", "right"), 6: (r["pb"], "0.00", "right"),
        7: (round(src["roe"], 4), pct1, "right"),
        8: (round(src["dy"], 4), pct2, "right"),
        9: (round(payout, 4), pct1, "right"),
        10: (round(src["g"], 4), pct1, "right"),
        11: (round(src["ret"], 4), pct1, "right"),
        12: ("● 持仓" if r["held"] else "观察", None, "center"),
        13: (r["note"], None, "left"),
    }
    for ci, (v, fmt, align) in vals.items():
        c = ws.cell(row=row, column=ci, value=v)
        c.border = thinb
        if fmt: c.number_format = fmt
        c.alignment = Alignment(horizontal=align, vertical="center", wrapText=(ci == 13))
        if ci in (5, 6, 8): c.fill = PatternFill("solid", fgColor=xl_color("#D9E2F3"))
        elif ci in (7, 9, 10, 11): c.fill = PatternFill("solid", fgColor=xl_color("#F2F2F2"))

ws.cell(row=AVG_ROW, column=2, value="观察池平均（简单算术）")
for ci, key in ((7, "roe"), (8, "dy"), (10, "g"), (11, "ret")):
    ok = [src[key] for src in raw.values() if src.get(key) is not None]
    c = ws.cell(row=AVG_ROW, column=ci, value=round(sum(ok) / len(ok), 4))
    c.number_format = pct2 if key == "dy" else pct1
for ci in range(1, 14):
    c = ws.cell(row=AVG_ROW, column=ci)
    c.font = Font(bold=True); c.border = thinb
    c.fill = PatternFill("solid", fgColor=xl_color("#D9E2F3"))

rng = f"K{DATA_START}:K{DATA_END}"
ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["0.15"], fill=PatternFill("solid", fgColor=xl_color("#C6EFCE")), font=Font(color=xl_color("#006100"), bold=True)))
ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0.08"], fill=PatternFill("solid", fgColor=xl_color("#FFC7CE")), font=Font(color=xl_color("#9C0006"))))

for col, w in {"A": 6, "B": 13, "C": 7, "D": 9.5, "E": 9, "F": 7.5, "G": 8.5, "H": 9, "I": 8.5, "J": 11, "K": 10.5, "L": 8.5, "M": 42}.items():
    ws.column_dimensions[col].width = w
ws.freeze_panes = "D3"
ws.auto_filter.ref = f"A2:M{DATA_END}"

ws2 = wb.create_sheet("口径说明")
notes = [
    ("公式", "理论回报 = 股息率 + 内生增长g；g = ROE × (1-分红率)；分红率 = 股息率 × PE（>100%按100%截断）；ROE = PB ÷ PE(TTM)（同源恒等式）。本表为静态快照，计算列已写入数值"),
    ("数据源", "腾讯行情源 gtimg（2026-10-16 收盘）：现价、PE(TTM)、PB、股息率（A股为行情字段直取，港股=每股DPS÷现价）"),
    ("覆盖", "投资平台观察池 31 只标的（含 7 只已持仓），无遗漏"),
    ("理论前提", "① PB 买卖时点不变；② ROE 可持续；③ 留存利润能赚出与存量相同的 ROE。任一崩塌则实际回报低于理论值——价值陷阱的来源"),
    ("† 泡泡玛特", "TTM ROE 53.2% 为利润率跳升年份的冲高值，g=39.5% 是把冲高值当永续外推的产物，实际可持续回报远低于 41.4%"),
    ("? 携程-S", "PE 5.4 疑似行情源字段错位（真实约 17-18），整行存疑待核"),
    ("‡ A/H 双上市", "H 股 ROE 受本币折算影响，与 A 股有 1-3 个点噪声（如平安 A 15.8% vs H 12.7%），H 股数字仅供参考"),
    ("对照", "书例（10% ROE/50%分红/2倍PB）= 7.5% 回报，约为本观察池中位水平；持仓 7 只全部在 11.7%-16.7% 区间"),
]
ws2.cell(row=1, column=1, value="口径说明").font = Font(bold=True, size=13)
for i, (k, v) in enumerate(notes, 3):
    a = ws2.cell(row=i, column=1, value=k); a.font = Font(bold=True); a.alignment = Alignment(vertical="top")
    b = ws2.cell(row=i, column=2, value=v); b.alignment = Alignment(wrapText=True, vertical="top")
    ws2.row_dimensions[i].height = 30
ws2.column_dimensions["A"].width = 12
ws2.column_dimensions["B"].width = 100

out = "/Users/jimislg/Documents/郎刚的芒格/codex-产出/观察池Gordon理论回报测算-20261017.xlsx"
wb.save(out)
print("saved:", out)
